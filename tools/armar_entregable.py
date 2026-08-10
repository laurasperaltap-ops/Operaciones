#!/usr/bin/env python3
"""Arma el libro entregable a partir del checkpoint de un anio.

Separa lo que ya esta resuelto de lo que pide intervencion humana, para que el
archivo se pueda accionar sin filtrar noventa columnas a mano.

Uso:
    python tools/armar_entregable.py --anio 2026
"""

from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from geocodificar import (COL_ANIO, COL_CIUDAD, COL_DESTINO, COL_ID, COL_ORIGEN,
                          DIR_CHECKPOINTS, FUENTE_POR_DEFECTO, cargar_procesados,
                          construir_salida)

ENCABEZADO = PatternFill("solid", fgColor="1F3864")
REVISAR = PatternFill("solid", fgColor="FFF2CC")
SIN_MATCH = PatternFill("solid", fgColor="FCE4E4")


def hojas_de_trabajo(lote: pd.DataFrame, procesados: dict) -> dict[str, pd.DataFrame]:
    """Divide el lote en la vista completa y las dos colas accionables."""
    completo = construir_salida(lote, procesados)

    # Vista compacta: lo minimo para decidir sobre una fila sin abrir el resto.
    columnas = [COL_ID, COL_CIUDAD, COL_ORIGEN, COL_DESTINO]
    for prefijo in ("Origen", "Destino"):
        columnas += [f"{prefijo} · Estado match", f"{prefijo} · Precisión",
                     f"{prefijo} · Estado", f"{prefijo} · Municipio",
                     f"{prefijo} · Parroquia", f"{prefijo} · Ciudad geocodificada",
                     f"{prefijo} · Dirección normalizada", f"{prefijo} · Motivo sin match"]
    compacta = completo[columnas]

    marca_o = completo["Origen · Estado match"]
    marca_d = completo["Destino · Estado match"]
    revisar = compacta[(marca_o == "Revisar") | (marca_d == "Revisar")]
    sin_match = compacta[(marca_o == "Sin match") | (marca_d == "Sin match")]

    return {"Direcciones 2026": completo,
            "Requieren revisión": revisar,
            "Sin match": sin_match}


def tabla_resumen(procesados: dict, total: int) -> pd.DataFrame:
    ETIQUETAS = {
        "exacta": "Match exacto", "aproximada": "Match aproximado",
        "nivel_ciudad": "Match a nivel ciudad",
        "coordenada_sin_nombre": "REVISAR · coordenada sin nombre",
        "conflicto_geografico": "REVISAR · fuera del área declarada",
        "solo_ciudad": "Sin match · no halló el lugar",
        "solo_pais": "Sin match · resolvió al país",
        "texto_vacio": "Sin match · celda vacía en la fuente",
        "sin_match": "Sin match · sin resultados",
        "sin_componentes": "Sin match · respuesta incompleta",
        "fuera_venezuela": "Sin match · fuera de Venezuela",
    }
    filas = []
    for clave, etiqueta in ETIQUETAS.items():
        origen = sum(1 for r in procesados.values() if r["origen"]["precision"] == clave)
        destino = sum(1 for r in procesados.values() if r["destino"]["precision"] == clave)
        if origen or destino:
            filas.append({"Resultado": etiqueta, "Origen": origen, "Destino": destino})

    resumen = pd.DataFrame(filas)

    totales = []
    for lado in ("origen", "destino"):
        m = sum(1 for r in procesados.values() if r[lado]["estado_match"] == "Match")
        rev = sum(1 for r in procesados.values() if r[lado]["estado_match"] == "Revisar")
        totales.append((m, rev, len(procesados) - m - rev))
    resumen.loc[len(resumen)] = {"Resultado": "— TOTAL con match —",
                                 "Origen": totales[0][0], "Destino": totales[1][0]}
    resumen.loc[len(resumen)] = {"Resultado": "— TOTAL a revisar —",
                                 "Origen": totales[0][1], "Destino": totales[1][1]}
    resumen.loc[len(resumen)] = {"Resultado": "— TOTAL sin match —",
                                 "Origen": totales[0][2], "Destino": totales[1][2]}
    ambos = sum(1 for r in procesados.values()
                if r["origen"]["estado_match"] == "Match"
                and r["destino"]["estado_match"] == "Match")
    resumen.loc[len(resumen)] = {"Resultado": f"Filas con AMBOS extremos resueltos (de {total})",
                                 "Origen": ambos, "Destino": ""}
    return resumen


def dar_formato(ruta: Path) -> None:
    """Encabezados fijos, filtros y ancho de columna legible en cada hoja."""
    import openpyxl

    libro = openpyxl.load_workbook(ruta)
    for hoja in libro.worksheets:
        if hoja.max_row < 1:
            continue
        for celda in hoja[1]:
            celda.fill = ENCABEZADO
            celda.font = Font(color="FFFFFF", bold=True, size=10)
            celda.alignment = Alignment(vertical="center", wrap_text=True)
        hoja.row_dimensions[1].height = 30
        hoja.freeze_panes = "A2"
        if hoja.max_row > 1:
            hoja.auto_filter.ref = hoja.dimensions

        for indice, columna in enumerate(hoja.iter_cols(), start=1):
            largos = [len(str(c.value)) for c in columna[:200] if c.value is not None]
            ancho = min(max(largos or [10]) + 2, 42)
            hoja.column_dimensions[get_column_letter(indice)].width = max(ancho, 11)

        # Colorea la fila segun necesite revision o no haya resuelto.
        encabezados = [c.value for c in hoja[1]]
        if "Origen · Estado match" in encabezados:
            cols = [encabezados.index(f"{p} · Estado match") for p in ("Origen", "Destino")]
            for fila in hoja.iter_rows(min_row=2):
                marcas = {fila[i].value for i in cols}
                relleno = REVISAR if "Revisar" in marcas else (
                    SIN_MATCH if "Sin match" in marcas else None)
                if relleno:
                    for celda in fila:
                        celda.fill = relleno
    libro.save(ruta)


def main() -> None:
    analizador = argparse.ArgumentParser(description=__doc__)
    analizador.add_argument("--anio", type=int, default=2026)
    analizador.add_argument("--fuente", type=Path, default=FUENTE_POR_DEFECTO)
    analizador.add_argument("--salida", type=Path, default=None)
    args = analizador.parse_args()

    marco = pd.read_excel(args.fuente, sheet_name=0).dropna(how="all")
    lote = marco[pd.to_numeric(marco[COL_ANIO], errors="coerce") == args.anio].copy()
    procesados = cargar_procesados(DIR_CHECKPOINTS / f"{args.anio}.jsonl")
    if not procesados:
        raise SystemExit(f"No hay checkpoint para {args.anio}.")

    salida = args.salida or Path("salidas") / f"Entregable_direcciones_{args.anio}.xlsx"
    salida.parent.mkdir(parents=True, exist_ok=True)

    hojas = hojas_de_trabajo(lote, procesados)
    with pd.ExcelWriter(salida, engine="openpyxl") as escritor:
        tabla_resumen(procesados, len(lote)).to_excel(escritor, sheet_name="Resumen", index=False)
        for nombre, marco_hoja in hojas.items():
            marco_hoja.to_excel(escritor, sheet_name=nombre[:31], index=False)

    escribir_excel_seguro(salida)
    dar_formato(salida)

    print(f"Entregable: {salida}")
    for nombre, marco_hoja in hojas.items():
        print(f"  hoja {nombre!r}: {len(marco_hoja)} filas × {len(marco_hoja.columns)} columnas")


def escribir_excel_seguro(ruta: Path) -> None:
    """Reaplica la preservacion de literales de error tras el ExcelWriter."""
    import openpyxl
    from openpyxl.cell.cell import TYPE_STRING

    libro = openpyxl.load_workbook(ruta)
    reparadas = 0
    for hoja in libro.worksheets:
        for fila in hoja.iter_rows():
            for celda in fila:
                if celda.data_type == "e" and celda.value is not None:
                    celda.data_type = TYPE_STRING
                    reparadas += 1
    if reparadas:
        libro.save(ruta)
        print(f"  ({reparadas} literales de error preservados como texto)")


if __name__ == "__main__":
    main()
